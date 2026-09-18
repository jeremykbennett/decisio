"""Tests for /api/campaigns/{id}/export/{type} xlsx export and Campaign Manager role label."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

CREDS = {
    "super": ("super@test.com", "newpass123"),
    "admin": ("admin@test.com", "admin123"),
    "marketer": ("marketer@test.com", "marketer123"),
}


def login(email, password):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json={"email": email, "password": password}, timeout=15)
    assert r.status_code == 200, f"login failed for {email}: {r.status_code} {r.text}"
    tok = r.json().get("access_token") or r.json().get("token")
    if tok:
        s.headers.update({"Authorization": f"Bearer {tok}"})
    return s


@pytest.fixture(scope="module")
def admin_session():
    return login(*CREDS["admin"])


@pytest.fixture(scope="module")
def marketer_session():
    return login(*CREDS["marketer"])


@pytest.fixture(scope="module")
def target_campaign(admin_session):
    r = admin_session.get(f"{BASE_URL}/api/campaigns?scope=all", timeout=15)
    assert r.status_code == 200
    camps = r.json()
    # Prefer 'Global Winter Push'
    for c in camps:
        if c.get("campaign_name") == "Global Winter Push":
            return c
    assert camps, "No campaigns available"
    return camps[0]


def _parse_xlsx(resp):
    assert resp.status_code == 200, f"status {resp.status_code}: {resp.text[:200]}"
    assert XLSX_MIME in resp.headers.get("content-type", ""), resp.headers.get("content-type")
    cd = resp.headers.get("content-disposition", "")
    assert ".xlsx" in cd, cd
    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    return rows


EXPECTED_HEADERS = ("Client Name", "Policy ID", "Platform", "Decision", "Updated By", "Updated At")


class TestExportBackend:
    def test_export_optin_admin(self, admin_session, target_campaign):
        r = admin_session.get(f"{BASE_URL}/api/campaigns/{target_campaign['id']}/export/opt_in", timeout=20)
        rows = _parse_xlsx(r)
        assert tuple(rows[0]) == EXPECTED_HEADERS
        for row in rows[1:]:
            assert row[3] == "Opt-in"

    def test_export_optout_admin(self, admin_session, target_campaign):
        r = admin_session.get(f"{BASE_URL}/api/campaigns/{target_campaign['id']}/export/opt_out", timeout=20)
        rows = _parse_xlsx(r)
        assert tuple(rows[0]) == EXPECTED_HEADERS
        for row in rows[1:]:
            assert row[3] == "Opt-out"

    def test_export_both_admin(self, admin_session, target_campaign):
        r = admin_session.get(f"{BASE_URL}/api/campaigns/{target_campaign['id']}/export/both", timeout=20)
        rows = _parse_xlsx(r)
        assert tuple(rows[0]) == EXPECTED_HEADERS
        decisions = {row[3] for row in rows[1:]}
        assert decisions.issubset({"Opt-in", "Opt-out"})

    def test_export_marketer_allowed(self, marketer_session, target_campaign):
        r = marketer_session.get(f"{BASE_URL}/api/campaigns/{target_campaign['id']}/export/both", timeout=20)
        rows = _parse_xlsx(r)
        assert tuple(rows[0]) == EXPECTED_HEADERS

    def test_export_invalid_type(self, admin_session, target_campaign):
        r = admin_session.get(f"{BASE_URL}/api/campaigns/{target_campaign['id']}/export/foo", timeout=15)
        assert r.status_code == 400

    def test_export_unknown_campaign(self, admin_session):
        r = admin_session.get(f"{BASE_URL}/api/campaigns/does-not-exist-xyz/export/both", timeout=15)
        assert r.status_code == 404


class TestGlobalWinterPushCounts:
    """Per seed context: Global Winter Push has 1 opt-in, 1 opt-out."""

    def test_specific_counts(self, admin_session):
        r = admin_session.get(f"{BASE_URL}/api/campaigns?scope=all", timeout=15)
        camp = next((c for c in r.json() if c.get("campaign_name") == "Global Winter Push"), None)
        if not camp:
            pytest.skip("Global Winter Push not present")

        oi = _parse_xlsx(admin_session.get(f"{BASE_URL}/api/campaigns/{camp['id']}/export/opt_in", timeout=20))
        oo = _parse_xlsx(admin_session.get(f"{BASE_URL}/api/campaigns/{camp['id']}/export/opt_out", timeout=20))
        both = _parse_xlsx(admin_session.get(f"{BASE_URL}/api/campaigns/{camp['id']}/export/both", timeout=20))

        oi_data = oi[1:]
        oo_data = oo[1:]
        both_data = both[1:]

        assert len(oi_data) == 1, f"expected 1 opt-in row, got {len(oi_data)}: {oi_data}"
        assert len(oo_data) == 1, f"expected 1 opt-out row, got {len(oo_data)}: {oo_data}"
        assert len(both_data) == 2

        # Client name checks
        assert oi_data[0][0] == "Meridian Benefits"
        assert oo_data[0][0] == "Surest Insurance Group"


class TestRoleStillMarketer:
    """Ensure underlying role key is still 'marketer' even after label rename."""

    def test_marketer_user_has_role_marketer(self):
        s = login(*CREDS["super"])
        r = s.get(f"{BASE_URL}/api/users", timeout=15)
        assert r.status_code == 200
        users = r.json()
        marketer = next((u for u in users if u.get("email") == "marketer@test.com"), None)
        assert marketer is not None
        assert marketer.get("role") == "marketer"
